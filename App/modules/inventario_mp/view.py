from __future__ import annotations

import getpass
import tkinter as tk
import os
import tempfile

from datetime import datetime
from decimal import Decimal
from tkinter import filedialog, messagebox, ttk

import customtkinter as ctk

from modules._shared.almox_repository import (
    AlmoxRepository,
    CampoInvalidoError,
)
from modules._shared.teclado_fixo import TecladoFixo
from modules.configuracoes.repository import ConfiguracoesRepository
from modules.inventario_mp.filtro_dialog import FiltroInventarioDialog


class InventarioMP(ctk.CTkFrame):
    """Lançamentos da contagem de MP."""

    CAMPOS = (
        ("bitola", "Material"),
        ("qualidade", "Qualidade"),
        ("rastreabilidade", "Rastreabilidade"),
        ("quantidade", "Quantidade"),
        ("localizacao", "Local"),
    )

    COLUNAS = (
        "data_hora",
        "bitola",
        "qualidade",
        "rastreabilidade",
        "quantidade",
        "localizacao",
        "operador",
    )

    COR_CAMPO_ATIVO = "#2F80ED"
    COR_ERRO = "#E74C3C"
    COR_SUCESSO = "#27AE60"
    COR_EDICAO = "#E6A23C"

    def __init__(self, parent) -> None:
        super().__init__(parent, corner_radius=0)

        self.repository: AlmoxRepository | None = None
        self.config_repository: ConfiguracoesRepository | None = None
        self.operador = getpass.getuser()
        self.registros_todos: dict[str, dict] = {}
        self.registros: dict[str, dict] = {}  # registros atualmente visíveis
        self.registro_em_edicao: dict | None = None
        self.campos: dict[str, object] = {}
        self.locais_estoque: list[str] = []
        self.filtros = {"data_de": None, "data_ate": None, "usuarios": None}
        self._indice_ativo = 0

        self.grid_columnconfigure(0, weight=0)
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self._build_formulario()
        self._build_tabela()

        self.after(100, self.atualizar)
        self.after(150, self._foco_inicial)

    # ------------------------------------------------------------------
    # Construção da tela
    # ------------------------------------------------------------------
    def _build_formulario(self) -> None:
        painel = ctk.CTkFrame(self, width=640)
        painel.grid(row=0, column=0, sticky="ns", padx=(20, 10), pady=20)
        painel.grid_propagate(False)
        painel.grid_columnconfigure(0, weight=1)
        self.painel = painel

        self.titulo_label = ctk.CTkLabel(
            painel,
            text="Inventário MP",
            anchor="w",
            font=ctk.CTkFont(size=25, weight="bold"),
        )
        self.titulo_label.grid(
            row=0,
            column=0,
            sticky="ew",
            padx=16,
            pady=(12, 4),
        )
        self._cor_titulo = self.titulo_label.cget("text_color")

        form = ctk.CTkFrame(painel, fg_color="transparent")
        form.grid(row=1, column=0, sticky="ew", padx=8)
        form.grid_columnconfigure(1, weight=1)

        for indice, (chave, rotulo) in enumerate(self.CAMPOS):
            ctk.CTkLabel(
                form,
                text=rotulo,
                width=140,
                anchor="w",
                font=ctk.CTkFont(size=16, weight="bold"),
            ).grid(row=indice, column=0, sticky="w", padx=(8, 8), pady=4)

            if chave == "localizacao":
                campo = ctk.CTkComboBox(
                    form,
                    values=[],
                    state="readonly",
                    height=35,
                    font=ctk.CTkFont(size=18),
                )
                campo.set("")
            else:
                campo = ctk.CTkEntry(
                    form,
                    height=35,
                    font=ctk.CTkFont(size=18),
                )
            campo.grid(row=indice, column=1, sticky="ew", padx=(0, 8), pady=4)
            campo.bind(
                "<FocusIn>",
                lambda _evento, i=indice: self._ao_focar_campo(i),
                add="+",
            )
            campo.bind(
                "<Return>",
                lambda _evento: self._proximo_campo(),
                add="+",
            )
            self.campos[chave] = campo

        self._borda_padrao = self.campos["bitola"].cget("border_color")

        botoes = ctk.CTkFrame(painel, fg_color="transparent")
        botoes.grid(row=2, column=0, sticky="ew", padx=8, pady=(8, 0))
        botoes.grid_columnconfigure((0, 1), weight=1, uniform="botoes")

        self.confirmar_btn = ctk.CTkButton(
            botoes,
            text="Incluir",
            height=50,
            font=ctk.CTkFont(size=18, weight="bold"),
            fg_color="#2E8B57",
            hover_color="#247447",
            command=self._confirmar,
        )
        self.confirmar_btn.grid(row=0, column=0, sticky="ew", padx=(8, 6))
        self.confirmar_btn.configure(state="disabled")

        self.limpar_btn = ctk.CTkButton(
            botoes,
            text="Limpar",
            height=50,
            font=ctk.CTkFont(size=18, weight="bold"),
            fg_color="#6C757D",
            hover_color="#5A6268",
            command=self._limpar,
        )
        self.limpar_btn.grid(row=0, column=1, sticky="ew", padx=(6, 8))

        self.status_label = ctk.CTkLabel(
            painel,
            text="",
            anchor="w",
            justify="left",
            wraplength=600,
            font=ctk.CTkFont(size=14),
        )
        self.status_label.grid(
            row=3,
            column=0,
            sticky="ew",
            padx=16,
            pady=(6, 0),
        )

        self.teclado = TecladoFixo(painel, ao_proximo=self._proximo_campo)
        self.teclado.grid(row=4, column=0, sticky="sew", padx=8, pady=(4, 10))
        painel.grid_rowconfigure(4, weight=1)

        for chave, campo in self.campos.items():
            if chave != "localizacao":
                self.teclado.registrar_campo(campo)
        self.teclado.definir_alvo(self.campos["bitola"])

    def _build_tabela(self) -> None:
        style = ttk.Style()
        style.configure(
            "InventarioMP.Treeview",
            font=("Arial", 12),
            rowheight=36,
        )
        style.configure(
            "InventarioMP.Treeview.Heading",
            font=("Arial", 14, "bold"),
        )

        direita = ctk.CTkFrame(self, fg_color="transparent")
        direita.grid(row=0, column=1, sticky="nsew", padx=(10, 20), pady=20)
        direita.grid_columnconfigure(0, weight=1)
        direita.grid_rowconfigure(1, weight=1)

        topo = ctk.CTkFrame(direita, fg_color="transparent")
        topo.grid(row=0, column=0, sticky="ew", pady=(0, 8))

        ctk.CTkLabel(
            topo,
            text="Itens Lançados",
            font=ctk.CTkFont(size=22, weight="bold"),
        ).pack(side="left")

        ctk.CTkButton(
            topo,
            text="Excluir",
            width=80,
            height=40,
            fg_color="#C0392B",
            hover_color="#922B21",
            command=self.excluir_selecionado,
        ).pack(side="right", padx=(0, 10))

        ctk.CTkButton(
            topo,
            text="Editar",
            width=50,
            height=40,
            command=self.editar_selecionado,
        ).pack(side="right", padx=(0, 10))

        self.contador_label = ctk.CTkLabel(topo, text="0 registro(s)")
        self.contador_label.pack(side="right", padx=(0, 16))

        container = ctk.CTkFrame(direita)
        container.grid(row=1, column=0, sticky="nsew")
        container.grid_columnconfigure(0, weight=1)
        container.grid_rowconfigure(0, weight=1)

        self.tree = ttk.Treeview(
            container,
            columns=self.COLUNAS,
            show="headings",
            selectmode="browse",
            style="InventarioMP.Treeview",
        )
        self.tree.tag_configure(
            "linha_par",
            background="#BEBEBE",
            foreground="#000000",
        )
        self.tree.tag_configure(
            "linha_impar",
            background="#FFFFFF",
            foreground="#000000",
        )

        rotulos = {
            "data_hora": "Dt/Hr",
            "bitola": "Material",
            "qualidade": "Qualidade",
            "rastreabilidade": "Rastreabilidade",
            "quantidade": "Quantidade",
            "localizacao": "Local",
            "operador": "Usuário",
        }
        larguras = {
            "data_hora": 140,
            "bitola": 90,
            "qualidade": 140,
            "rastreabilidade": 155,
            "quantidade": 115,
            "localizacao": 130,
            "operador": 140,
        }
        for coluna in self.COLUNAS:
            self.tree.heading(coluna, text=rotulos[coluna])
            self.tree.column(
                coluna,
                width=larguras[coluna],
                minwidth=larguras[coluna],
                stretch=False,
                anchor="center",
            )
        for coluna in ("qualidade", "rastreabilidade", "operador"):
            self.tree.column(coluna, anchor="w")

        self.tree.bind("<Double-1>", lambda _evento: self.editar_selecionado())

        y_scroll = ctk.CTkScrollbar(
            container,
            orientation="vertical",
            command=self.tree.yview,
            width=22,
            fg_color=("gray85", "gray20"),
            button_color=("#557A95", "#557A95"),
            button_hover_color=("#2F80ED", "#2F80ED"),
            corner_radius=6,
            border_spacing=3,
        )
        x_scroll = ctk.CTkScrollbar(
            container,
            orientation="horizontal",
            command=self.tree.xview,
            height=22,
            fg_color=("gray85", "gray20"),
            button_color=("#557A95", "#557A95"),
            button_hover_color=("#2F80ED", "#2F80ED"),
            corner_radius=6,
            border_spacing=3,
        )
        self.tree.configure(
            yscrollcommand=y_scroll.set,
            xscrollcommand=x_scroll.set,
        )
        self.tree.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")


        # Barra de ações abaixo da tabela, mantendo o espaço principal para os registros.
        rodape = ctk.CTkFrame(direita, fg_color="transparent")
        rodape.grid(row=2, column=0, sticky="ew", pady=(8, 0))

        ctk.CTkButton(
            rodape,
            text="Filtros",
            width=100,
            height=38,
            command=self.abrir_filtros,
        ).pack(side="right")
        ctk.CTkButton(
            rodape,
            text="Excel",
            width=90,
            height=38,
            fg_color="#217346",
            hover_color="#185C37",
            command=self.exportar_excel,
        ).pack(side="right", padx=(0, 8))
        ctk.CTkButton(
            rodape,
            text="Atualizar",
            width=100,
            height=38,
            command=self.atualizar,
        ).pack(side="right", padx=(0, 8))

    # ------------------------------------------------------------------
    # Foco e navegação entre os campos
    # ------------------------------------------------------------------
    def _foco_inicial(self) -> None:
        try:
            if self.winfo_exists():
                self.campos["bitola"].focus_set()
        except tk.TclError:
            pass

    def _ao_focar_campo(self, indice: int) -> None:
        self._indice_ativo = indice
        for i, (chave, _rotulo) in enumerate(self.CAMPOS):
            self.campos[chave].configure(
                border_color=(
                    self.COR_CAMPO_ATIVO if i == indice else self._borda_padrao
                )
            )
        if self.CAMPOS[indice][0] == "localizacao":
            self.teclado.definir_alvo(None)

    def _proximo_campo(self) -> None:
        proximo = (self._indice_ativo + 1) % len(self.CAMPOS)
        self.campos[self.CAMPOS[proximo][0]].focus_set()

    # ------------------------------------------------------------------
    # Dados
    # ------------------------------------------------------------------
    def _obter_repositorio(self) -> AlmoxRepository:
        if self.repository is None:
            self.repository = AlmoxRepository()
        return self.repository

    def atualizar(self, selecionar_id=None) -> None:
        try:
            if not self.winfo_exists():
                return
        except tk.TclError:
            return

        try:
            self._atualizar_locais_estoque()
        except Exception as exc:
            self.locais_estoque = []
            self.campos["localizacao"].configure(values=[])
            self.campos["localizacao"].set("")
            self.confirmar_btn.configure(state="disabled")
            self._mostrar_status(
                f"Não foi possível carregar os locais de estoque: {exc}",
                erro=True,
            )

        try:
            registros = self._obter_repositorio().listar_inventario_mp()
        except Exception as exc:
            messagebox.showerror(
                "Inventário MP",
                str(exc),
                parent=self.winfo_toplevel(),
            )
            return

        self._preencher_tabela(registros, selecionar_id)

    
    def exportar_excel(self) -> None:
        """Exporta os registros carregados para uma planilha Excel."""
        if not self.registros:
            messagebox.showinfo(
                "Exportar Excel",
                "Não existem registros para exportar.",
                parent=self.winfo_toplevel(),
            )
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import (
                Font,
                PatternFill,
                Alignment,
                Border,
                Side,
            )

            wb = Workbook()
            ws = wb.active
            ws.title = "Inventário MP"

            cabecalhos = [
                "Dt/Hr",
                "Material",
                "Qualidade",
                "Rastreabilidade",
                "Quantidade",
                "Local",
                "Usuário",
            ]

            ws.append(cabecalhos)

            # Preenche os registros mantendo a ordem da tabela.
            for row in self.registros.values():

                quantidade = row.get("quantidade")
                quantidade = (
                    float(Decimal(str(quantidade)))
                    if quantidade not in (None, "")
                    else 0.0
                )

                ws.append([
                    self._data_excel(row.get("criado_em")),
                    row.get("bitola") or "",
                    row.get("qualidade") or "",
                    row.get("rastreabilidade") or "",
                    quantidade,
                    row.get("localizacao") or "",
                    row.get("usuario") or "",
                ])

            # Formato de data e hora.
            for celula in ws["A"][1:]:
                celula.number_format = "dd/mm/yyyy hh:mm"

            # Formato numérico da quantidade.
            for celula in ws["E"][1:]:
                celula.number_format = "#,##0.00"

            # Largura das colunas.
            larguras = {
                "A": 20,
                "B": 18,
                "C": 15,
                "D": 20,
                "E": 16,
                "F": 10,
                "G": 15,
            }

            for coluna, largura in larguras.items():
                ws.column_dimensions[coluna].width = largura

            # Melhorias de navegação.
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            # Cria um arquivo temporário no Windows.
            fd, caminho = tempfile.mkstemp(
                prefix=f"inventario_mp_{datetime.now():%Y%m%d_%H%M%S}_",
                suffix=".xlsx",
            )

            os.close(fd)

            # Gera a planilha.
            wb.save(caminho)

            # Abre automaticamente no aplicativo padrão do Windows.
            os.startfile(caminho)

        except ImportError:
            messagebox.showerror(
                "Biblioteca ausente",
                "Instale a biblioteca openpyxl:\n\n"
                "pip install openpyxl",
                parent=self.winfo_toplevel(),
            )

        except Exception as exc:
            messagebox.showerror(
                "Erro na exportação",
                str(exc),
                parent=self.winfo_toplevel(),
            )

    def _preencher_tabela(self, registros: list[dict], selecionar_id=None) -> None:
        self.registros_todos = {
            str(row["id"]): row
            for row in registros
        }
        self._aplicar_filtros_tabela(selecionar_id)

    def _aplicar_filtros_tabela(self, selecionar_id=None) -> None:
        data_de = self.filtros.get("data_de")
        data_ate = self.filtros.get("data_ate")
        usuarios = self.filtros.get("usuarios")

        filtrados: dict[str, dict] = {}
        for registro_id, row in self.registros_todos.items():
            usuario = str(row.get("usuario") or "(sem usuário)")
            if usuarios is not None and usuario not in usuarios:
                continue

            data_registro = self._data_local_registro(row.get("criado_em"))
            if data_de and (data_registro is None or data_registro < data_de):
                continue
            if data_ate and (data_registro is None or data_registro > data_ate):
                continue

            filtrados[registro_id] = row

        self.registros = filtrados
        for item_id in self.tree.get_children():
            self.tree.delete(item_id)

        for indice, (chave, row) in enumerate(self.registros.items()):
            tag = "linha_par" if indice % 2 == 0 else "linha_impar"
            self.tree.insert(
                "",
                "end",
                iid=chave,
                values=(
                    self._fmt_data_hora(row.get("criado_em")),
                    row.get("bitola") or "",
                    row.get("qualidade") or "",
                    row.get("rastreabilidade") or "",
                    self._fmt(row.get("quantidade")),
                    row.get("localizacao") or "",
                    row.get("usuario") or "",
                ),
                tags=(tag,),
            )

        total = len(self.registros_todos)
        visiveis = len(self.registros)
        if self.filtros.get("data_de") or self.filtros.get("data_ate") or usuarios is not None:
            self.contador_label.configure(text=f"{visiveis} de {total} registro(s)")
        else:
            self.contador_label.configure(text=f"{total} registro(s)")

        if selecionar_id is not None and str(selecionar_id) in self.registros:
            self.tree.selection_set(str(selecionar_id))
            self.tree.see(str(selecionar_id))

    def _data_local_registro(self, valor):
        if not valor:
            return None
        try:
            data_hora = datetime.fromisoformat(
                str(valor).strip().replace("Z", "+00:00")
            )
        except ValueError:
            return None
        if data_hora.tzinfo is not None:
            data_hora = data_hora.astimezone()
        return data_hora.date()

    def _obter_config_repository(self) -> ConfiguracoesRepository:
        if self.config_repository is None:
            self.config_repository = ConfiguracoesRepository()
        return self.config_repository

    def _atualizar_locais_estoque(self) -> None:
        locais = self._obter_config_repository().listar_locais_estoque(somente_ativos=True)
        self.locais_estoque = [
            str(item.get("nome") or "").strip()
            for item in locais
            if str(item.get("nome") or "").strip()
        ]

        valores = list(self.locais_estoque)
        # Um local bloqueado pode permanecer selecionado apenas no registro que
        # já o utilizava. O usuário pode editar os demais campos sem migrá-lo.
        if self.registro_em_edicao is not None:
            local_anterior = str(self.registro_em_edicao.get("localizacao") or "").strip()
            if local_anterior and local_anterior not in valores:
                valores.append(local_anterior)

        campo = self.campos["localizacao"]
        atual = campo.get().strip()
        campo.configure(values=valores)
        campo.set(atual if atual in valores else "")
        self.confirmar_btn.configure(
            state="normal" if self.locais_estoque or self.registro_em_edicao else "disabled"
        )

        if not self.locais_estoque:
            self._mostrar_status(
                "Não existem locais ativos. Cadastre ou libere um local em Configurações.",
                erro=True,
            )

    def abrir_filtros(self) -> None:
        usuarios = sorted({
            str(row.get("usuario") or "(sem usuário)")
            for row in self.registros_todos.values()
        }, key=str.casefold)
        FiltroInventarioDialog(
            self,
            usuarios=usuarios,
            filtros_atuais=self.filtros,
            ao_aplicar=self._receber_filtros,
            ao_limpar=self._limpar_filtros,
        )

    def _receber_filtros(self, filtros: dict) -> None:
        self.filtros = filtros
        self._aplicar_filtros_tabela()

    def _limpar_filtros(self) -> None:
        self.filtros = {"data_de": None, "data_ate": None, "usuarios": None}
        self._aplicar_filtros_tabela()

    # ------------------------------------------------------------------
    # Incluir / editar / excluir / limpar
    # ------------------------------------------------------------------
    def _ler_campos(self) -> dict[str, str]:
        return {
            chave: self.campos[chave].get().strip()
            for chave, _rotulo in self.CAMPOS
        }

    def _confirmar(self) -> None:
        valores = self._ler_campos()
        editando = self.registro_em_edicao

        try:
            repositorio = self._obter_repositorio()
            if editando is None:
                resultado = repositorio.incluir_inventario_mp(
                    **valores,
                    nome_operador=self.operador,
                )
                mensagem = "Registro incluído com sucesso."
            else:
                resultado = repositorio.atualizar_inventario_mp(
                    int(editando["id"]),
                    **valores,
                    nome_operador=self.operador,
                )
                mensagem = "Registro atualizado com sucesso."
        except CampoInvalidoError as exc:
            self._mostrar_status(str(exc), erro=True)
            self.campos[exc.campo].focus_set()
            return
        except Exception as exc:
            messagebox.showerror(
                "Inventário MP",
                str(exc),
                parent=self.winfo_toplevel(),
            )
            return

        registro_id = resultado.get("id")
        self._limpar(manter_status=True)
        self._mostrar_status(mensagem)
        self.atualizar(selecionar_id=registro_id)

    def editar_selecionado(self) -> None:
        selecionados = self.tree.selection()
        if not selecionados:
            messagebox.showinfo(
                "Inventário MP",
                "Selecione um registro na tabela.",
                parent=self.winfo_toplevel(),
            )
            return

        row = self.registros.get(selecionados[0])
        if not row:
            self.atualizar()
            return

        for chave, _rotulo in self.CAMPOS:
            campo = self.campos[chave]
            if chave == "localizacao":
                local_atual = str(row.get(chave) or "")
                opcoes = list(self.locais_estoque)
                if local_atual and local_atual not in opcoes:
                    opcoes.append(local_atual)
                campo.configure(values=opcoes)
                campo.set(local_atual)
                continue
            campo.delete(0, "end")
            if chave == "quantidade":
                texto = self._fmt(row.get("quantidade")).replace(".", ",")
            else:
                texto = str(row.get(chave) or "")
            campo.insert(0, texto)

        self.registro_em_edicao = row
        self.titulo_label.configure(
            text=f"Editando registro nº {row['id']}",
            text_color=self.COR_EDICAO,
        )
        self.confirmar_btn.configure(
            text="Salvar alteração",
            fg_color="#2F80ED",
            hover_color="#1F6AA5",
        )
        self.limpar_btn.configure(text="Cancelar edição")
        self.confirmar_btn.configure(state="normal")
        self._mostrar_status("")
        self.campos["bitola"].focus_set()

    def excluir_selecionado(self) -> None:
        selecionados = self.tree.selection()
        if not selecionados:
            messagebox.showinfo(
                "Inventário MP",
                "Selecione um registro na tabela para excluir.",
                parent=self.winfo_toplevel(),
            )
            return

        registro_id = selecionados[0]
        row = self.registros.get(registro_id)
        if not row:
            self.atualizar()
            return

        material = row.get("bitola") or "(sem material)"
        rastreabilidade = row.get("rastreabilidade") or "(sem rastreabilidade)"
        confirmado = messagebox.askyesno(
            "Confirmar exclusão",
            (
                f"Deseja realmente excluir o registro nº {registro_id}?\n\n"
                f"Material: {material}\n"
                f"Rastreabilidade: {rastreabilidade}\n\n"
                "Esta ação não poderá ser desfeita."
            ),
            icon="warning",
            parent=self.winfo_toplevel(),
        )
        if not confirmado:
            return

        try:
            # A operação é feita por RPC, seguindo o mesmo padrão usado
            # para inclusão e alteração do inventário. A função do banco
            # valida o acesso do usuário ao módulo antes de excluir.
            response = self._obter_repositorio().client.rpc(
                "excluir_inventario_mp",
                {
                    "p_id": int(registro_id),
                    "p_usuario": self.operador,
                },
            ).execute()
            resultado = response.data
            if isinstance(resultado, list) and resultado:
                resultado = resultado[0]
            if not isinstance(resultado, dict) or resultado.get("status") != "EXCLUIDO":
                raise RuntimeError(
                    "O Supabase não confirmou a exclusão do registro."
                )
        except Exception as exc:
            messagebox.showerror(
                "Falha ao excluir registro",
                str(exc),
                parent=self.winfo_toplevel(),
            )
            return

        if (
            self.registro_em_edicao is not None
            and str(self.registro_em_edicao.get("id")) == str(registro_id)
        ):
            self._limpar(manter_status=True)

        self._mostrar_status(f"Registro nº {registro_id} excluído com sucesso.")
        self.atualizar()

    def _limpar(self, manter_status: bool = False) -> None:
        for chave, campo in self.campos.items():
            if chave == "localizacao":
                campo.set("")
                campo.configure(values=self.locais_estoque)
            else:
                campo.delete(0, "end")

        self.registro_em_edicao = None
        self.titulo_label.configure(
            text="Inventário MP",
            text_color=self._cor_titulo,
        )
        self.confirmar_btn.configure(
            text="Incluir",
            fg_color="#2E8B57",
            hover_color="#247447",
        )
        self.limpar_btn.configure(text="Limpar")
        self.confirmar_btn.configure(
            state="normal" if self.locais_estoque else "disabled"
        )

        if not manter_status:
            self._mostrar_status("")
        self.campos["bitola"].focus_set()

    def _mostrar_status(self, texto: str, erro: bool = False) -> None:
        self.status_label.configure(
            text=texto,
            text_color=self.COR_ERRO if erro else self.COR_SUCESSO,
        )

    # ------------------------------------------------------------------
    # Formatação
    # ------------------------------------------------------------------
    @staticmethod
    def _fmt(valor) -> str:
        if valor in (None, ""):
            return "0"
        numero = Decimal(str(valor))
        return f"{numero:.3f}".rstrip("0").rstrip(".") or "0"

    @staticmethod
    def _fmt_data_hora(valor) -> str:
        if not valor:
            return ""
        try:
            data_hora = datetime.fromisoformat(
                str(valor).strip().replace("Z", "+00:00")
            )
        except ValueError:
            return str(valor)
        if data_hora.tzinfo is not None:
            data_hora = data_hora.astimezone()
        return data_hora.strftime("%d/%m/%Y %H:%M")

    @staticmethod
    def _data_excel(valor):
        if not valor:
            return None

        try:
            data_hora = datetime.fromisoformat(
                str(valor).strip().replace("Z", "+00:00")
            )
        except ValueError:
            return str(valor)

        if data_hora.tzinfo is not None:
            data_hora = data_hora.astimezone().replace(tzinfo=None)

        return data_hora
